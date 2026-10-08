"""Create simple, local PowerPoint decks from a declarative slide specification.

python-pptx is optional. This module does not fetch assets, interpret code, or run
model output. Callers provide already-reviewed title, slide text, and notes.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

MAX_SLIDES = 30
MAX_TITLE = 200
MAX_SLIDE_TITLE = 180
MAX_BODY = 5_000
MAX_BULLETS = 12
MAX_BULLET = 1_000
MAX_NOTES = 2_000
SPEC_GUIDE = '''Return one JSON object with this schema; do not include Markdown fences or extra keys:
{"title":"Short deck title","slides":[{"title":"Slide title","body":"Optional short paragraph","bullets":["Optional point"],"speaker_notes":"Optional presenter notes"}]}
Use 1 to 30 slides. Each slide supports title, optional body, optional bullets, and optional speaker_notes. Keep slide text concise. Do not include images, URLs to assets, code, or instructions to fetch external resources.'''


def _text(value: Any, label: str, limit: int, *, required: bool = False) -> str:
    if value is None and not required:
        return ""
    if not isinstance(value, str):
        raise ValueError(f"{label} must be text.")
    value = value.strip()
    if required and not value:
        raise ValueError(f"{label} cannot be empty.")
    if len(value) > limit:
        raise ValueError(f"{label} exceeds the {limit:,}-character limit.")
    return value


def _validate_spec(spec: Any) -> tuple[str, list[dict[str, Any]]]:
    if not isinstance(spec, dict) or set(spec) != {"title", "slides"}:
        raise ValueError("Presentation spec must contain only title and slides.")
    title = _text(spec["title"], "Presentation title", MAX_TITLE, required=True)
    slides = spec["slides"]
    if not isinstance(slides, list) or not 1 <= len(slides) <= MAX_SLIDES:
        raise ValueError(f"Presentation must contain between 1 and {MAX_SLIDES} slides.")
    clean: list[dict[str, Any]] = []
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict) or set(slide) - {"title", "body", "bullets", "speaker_notes"}:
            raise ValueError(f"Slide {index} may contain only title, body, bullets, and speaker_notes.")
        slide_title = _text(slide.get("title"), f"Slide {index} title", MAX_SLIDE_TITLE, required=True)
        body = _text(slide.get("body", ""), f"Slide {index} body", MAX_BODY)
        notes = _text(slide.get("speaker_notes", ""), f"Slide {index} speaker notes", MAX_NOTES)
        bullets = slide.get("bullets", [])
        if not isinstance(bullets, list) or len(bullets) > MAX_BULLETS:
            raise ValueError(f"Slide {index} may contain at most {MAX_BULLETS} bullet items.")
        clean_bullets = [_text(item, f"Slide {index} bullet", MAX_BULLET, required=True) for item in bullets]
        clean.append({"title": slide_title, "body": body, "bullets": clean_bullets, "notes": notes})
    return title, clean


def safe_output_basename(title: str) -> str:
    """Create a portable basename from a title; never accept path components."""
    if not isinstance(title, str):
        raise ValueError("Presentation title must be text.")
    slug = re.sub(r"[^A-Za-z0-9_-]+", "-", title.strip()).strip("-_")[:72]
    return slug or "presentation"


def _presentation_module():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as error:
        raise RuntimeError("PowerPoint creation needs the optional python-pptx package. Install it in Mavi's Python environment.") from error
    return Presentation, RGBColor, MSO_SHAPE, PP_ALIGN, Inches, Pt


def create_presentation(spec: dict[str, Any], output_dir: str | os.PathLike) -> Path:
    """Create a clean 16:9 monochrome PPTX without overwriting an existing file."""
    title, slides = _validate_spec(spec)
    raw_dir = Path(output_dir).expanduser()
    if raw_dir.is_symlink() or not raw_dir.is_dir():
        raise ValueError("Presentation output folder must be an existing, non-symlink directory.")
    folder = raw_dir.resolve(strict=True)
    if folder.is_symlink():
        raise ValueError("Presentation output folder must be an existing, non-symlink directory.")

    Presentation, RGBColor, MSO_SHAPE, PP_ALIGN, Inches, Pt = _presentation_module()
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = title
    presentation.core_properties.subject = "Created locally by Mavi"
    blank = presentation.slide_layouts[6]

    for number, item in enumerate(slides, 1):
        slide = presentation.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

        title_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.62), Inches(11.9), Inches(0.78))
        title_frame = title_box.text_frame
        title_frame.clear()
        title_frame.word_wrap = True
        paragraph = title_frame.paragraphs[0]
        paragraph.text = item["title"]
        paragraph.font.name = "Aptos Display"
        paragraph.font.size = Pt(28)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(20, 20, 20)

        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.48), Inches(11.9), Inches(0.025))
        rule.fill.solid()
        rule.fill.fore_color.rgb = RGBColor(190, 190, 190)
        rule.line.fill.background()

        content = slide.shapes.add_textbox(Inches(0.9), Inches(1.82), Inches(11.4), Inches(4.9))
        frame = content.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.04)
        frame.margin_right = Inches(0.04)
        frame.margin_top = Inches(0.04)
        frame.margin_bottom = Inches(0.04)
        first = True
        if item["body"]:
            paragraph = frame.paragraphs[0]
            paragraph.text = item["body"]
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(19)
            paragraph.font.color.rgb = RGBColor(55, 55, 55)
            first = False
        for bullet in item["bullets"]:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.text = "•\u00a0" + bullet
            paragraph.level = 0
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(18)
            paragraph.font.color.rgb = RGBColor(45, 45, 45)
            paragraph.space_after = Pt(13)
        if first:
            frame.paragraphs[0].text = ""

        footer = slide.shapes.add_textbox(Inches(11.85), Inches(7.08), Inches(0.7), Inches(0.2))
        footer_p = footer.text_frame.paragraphs[0]
        footer_p.text = f"{number:02d}"
        footer_p.alignment = PP_ALIGN.RIGHT
        footer_p.font.name = "Aptos"
        footer_p.font.size = Pt(9)
        footer_p.font.color.rgb = RGBColor(115, 115, 115)

        if item["notes"]:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is None:
                raise RuntimeError("This python-pptx version cannot create speaker notes.")
            notes_frame.text = item["notes"]

    base = safe_output_basename(title)
    counter = 1
    while True:
        suffix = "" if counter == 1 else f"-{counter}"
        target = folder / f"{base}{suffix}.pptx"
        if target.parent != folder or target.name in {".", ".."}:
            raise ValueError("Presentation output path is invalid.")
        try:
            with target.open("xb") as stream:
                try:
                    presentation.save(stream)
                    stream.flush()
                    os.fsync(stream.fileno())
                except Exception:
                    stream.close()
                    target.unlink(missing_ok=True)
                    raise
            return target
        except FileExistsError:
            counter += 1
