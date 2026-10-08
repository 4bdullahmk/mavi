import hashlib
import importlib.util
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("mavi_package_source", ROOT / "scripts/package_source.py")
package = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = package
SPEC.loader.exec_module(package)


class SourcePackageTests(unittest.TestCase):
    def test_windows_installer_is_packaged_with_explicit_safe_setup_controls(self):
        installer = (ROOT / "Install-Mavi.ps1").read_text(encoding="utf-8")
        command = (ROOT / "Install-Mavi.cmd").read_text(encoding="utf-8")
        self.assertIn('Python.Python.3.12', installer)
        self.assertIn('Ollama.Ollama', installer)
        self.assertIn("'--scope', 'user'", installer)
        self.assertNotIn("-ExecutionPolicy", command)
        self.assertNotIn("-ExecutionPolicy Bypass", installer + command)
        self.assertNotIn("--accept-package-agreements", installer)
        self.assertIn("-DownloadModel", installer)
        self.assertIn("-EnableWindowsAutomation", installer)
        self.assertIn("-EnableDictation", installer)
        self.assertIn("Test-PortClosed 8769", installer)
        self.assertIn("/api/chat", installer)
        self.assertIn("requirements-installed.txt", installer)
        self.assertIn("--spreadsheet-smoke", installer)
        self.assertIn("theme = 'system'", installer)
        self.assertIn("discord = [pscustomobject]@{}", installer)
        self.assertLess(installer.index("$installedNames = @($tags.models"), installer.index("$savedModel -in $installedNames"))
        self.assertIn('Install-Mavi.ps1', package._ROOT_FILES)
        self.assertIn('Install-Mavi.cmd', package._ROOT_FILES)

    def test_two_user_discord_workspaces_stay_private_across_verified_update(self):
        """Package and updater touch release files, never separate user workspaces."""
        import json

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source"
            source.mkdir()
            for name in package._ROOT_FILES:
                (source / name).write_text(f"fixture {name}\n", encoding="utf-8")
            app = source / "app"
            app.mkdir()
            required_helpers = (
                "FileTools.py", "DeveloperAgent.py", "ProjectReadCache.py", "CADAgent.py",
                "CADTools.py", "ProductSpecs.py", "MeshTools.py", "StockTools.py", "SelfUpdate.py",
                "SpreadsheetTools.py", "PresentationTools.py",
            )
            for name in required_helpers:
                (app / name).write_text("# packaged test helper\n", encoding="utf-8")
            portable = source / "portable"
            portable.mkdir()
            (portable / "server.py").write_text("# clean server fixture\n", encoding="utf-8")
            (portable / "requirements.txt").write_text("\n", encoding="utf-8")

            # Seed tempting-but-private fixture files in ignored source data paths.
            # These are synthetic values, never copied from a real installation.
            secret_marker = "synthetic-test-token-" + "NOT-A-REAL-CREDENTIAL"
            history_marker = "synthetic-private-chat-history-" + "NOT-REAL"
            for folder in (source / "data", portable / "data"):
                folder.mkdir()
                (folder / "discord.json").write_text(
                    json.dumps({"token": secret_marker, "channel_id": "111111111111111111"}),
                    encoding="utf-8",
                )
                (folder / "history.json").write_text(history_marker, encoding="utf-8")

            archive, sidecar = package.package_source(source, root / "artifacts", "2.0.0-isolation")
            digest = hashlib.sha256(archive.read_bytes()).hexdigest()
            self.assertEqual(sidecar.read_text(encoding="ascii").split()[0], digest)
            with zipfile.ZipFile(archive) as zipped:
                names = zipped.namelist()
                payload = b"\n".join(zipped.read(name) for name in names)
                self.assertFalse(any("/data/" in name or name.startswith("data/") for name in names))
                self.assertNotIn(secret_marker.encode(), payload)
                self.assertNotIn(history_marker.encode(), payload)

            update_spec = importlib.util.spec_from_file_location("mavi_update_isolation", ROOT / "Update-Mavi.py")
            update = importlib.util.module_from_spec(update_spec)
            sys.modules[update_spec.name] = update
            update_spec.loader.exec_module(update)

            workspaces = []
            for label, channel_id, user_id in (
                ("friend_a", "222222222222222222", "333333333333333333"),
                ("friend_b", "444444444444444444", "555555555555555555"),
            ):
                workspace = root / "users" / label / "Mavi"
                workspace.mkdir(parents=True)
                discord = {
                    "channel_id": channel_id,
                    "user_ids": [user_id],
                    "allow_tasks": False,
                    # The real runtime never persists this field. Its absence is
                    # part of the fixture's memory-only token contract.
                }
                history = {"chats": [{"id": label, "messages": [{"text": f"{label} private fixture"}]}]}
                (workspace / "discord.json").write_text(json.dumps(discord), encoding="utf-8")
                (workspace / "history.json").write_text(json.dumps(history), encoding="utf-8")
                workspaces.append((workspace, discord, history))

            releases = root / "users" / "releases"
            for fake_pid, (workspace, expected_discord, expected_history) in zip((101, 202), workspaces):
                with patch.object(update.os, "getpid", return_value=fake_pid):
                    launcher = update.install_verified_archive(archive, releases, digest)
                self.assertTrue((launcher.parent / "portable" / "server.py").is_file())
                self.assertEqual(json.loads((workspace / "discord.json").read_text(encoding="utf-8")), expected_discord)
                self.assertEqual(json.loads((workspace / "history.json").read_text(encoding="utf-8")), expected_history)
                self.assertNotIn("token", json.loads((workspace / "discord.json").read_text(encoding="utf-8")))

    def test_current_source_package_is_reproducible_and_update_extractable(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "release"
            first, sidecar = package.package_source(ROOT, out, "2.0.0-test")
            first_bytes = first.read_bytes()
            digest = hashlib.sha256(first_bytes).hexdigest()
            self.assertEqual(sidecar.read_text().split()[0], digest)
            with zipfile.ZipFile(first) as zf:
                names = zf.namelist()
                self.assertIn("Mavi-Windows-2.0.0-test/portable/server.py", names)
                self.assertIn("Mavi-Windows-2.0.0-test/portable/image_runtime.py", names)
                self.assertIn("Mavi-Windows-2.0.0-test/portable/model_policy.py", names)
                self.assertIn("Mavi-Windows-2.0.0-test/portable/web/mavi-mark.png", names)
                self.assertIn("Mavi-Windows-2.0.0-test/app/SpreadsheetTools.py", names)
                self.assertIn("Mavi-Windows-2.0.0-test/app/PresentationTools.py", names)
                for helper in ("FileTools.py", "DeveloperAgent.py", "ProjectReadCache.py", "CADAgent.py",
                               "CADTools.py", "ProductSpecs.py", "MeshTools.py", "StockTools.py", "SelfUpdate.py"):
                    self.assertIn("Mavi-Windows-2.0.0-test/app/" + helper, names)
                for support in ("Setup-Models.py", "AGENTS.md", "CHANGELOG.md"):
                    self.assertIn("Mavi-Windows-2.0.0-test/" + support, names)
                self.assertIn("Mavi-Windows-2.0.0-test/Install-Mavi.cmd", names)
                self.assertIn("Mavi-Windows-2.0.0-test/Install-Mavi.ps1", names)
                for support in ("INSTALL_WITH_CODEX.md", "Install-Mavi.cmd", "Install-Mavi.ps1", "test_update.py"):
                    self.assertIn("Mavi-Windows-2.0.0-test/" + support, names)
                self.assertIn("Mavi-Windows-2.0.0-test/tests/test_package_source.py", names)
                self.assertIn("Mavi-Windows-2.0.0-test/scripts/package_source.py", names)
                self.assertFalse(any("/data/" in n or "/logs/" in n or ".pyc" in n for n in names))
                self.assertFalse(any(Path(n).is_absolute() or ".." in Path(n).parts for n in names))
            second, _ = package.package_source(ROOT, out, "2.0.0-test")
            self.assertEqual(first_bytes, second.read_bytes())

            update_spec = importlib.util.spec_from_file_location("mavi_update", ROOT / "Update-Mavi.py")
            update = importlib.util.module_from_spec(update_spec)
            sys.modules[update_spec.name] = update
            update_spec.loader.exec_module(update)
            launcher = update.install_verified_archive(first, Path(tmp) / "installed", digest)
            self.assertEqual(launcher.name, "Start-Mavi.cmd")
            self.assertTrue((launcher.parent / "portable/server.py").is_file())
            self._exercise_extracted_package(launcher.parent)

    @staticmethod
    def _exercise_extracted_package(release_root):
        """Exercise the actual portable adapter and both generated-file builders."""
        import importlib
        import json
        import os
        import tempfile

        optional = importlib.util.find_spec("openpyxl") and importlib.util.find_spec("pptx")
        if not optional:
            raise unittest.SkipTest("Install the portable xlsx/pptx requirements to run artifact integration.")
        sys.path.insert(0, str(release_root / "portable"))
        module = importlib.util.spec_from_file_location("packaged_tools_runtime", release_root / "portable/tools_runtime.py")
        runtime = importlib.util.module_from_spec(module)
        sys.modules[module.name] = runtime
        module.loader.exec_module(runtime)
        for helper in ("FileTools", "DeveloperAgent", "ProjectReadCache", "CADAgent", "CADTools",
                       "ProductSpecs", "MeshTools", "StockTools", "SelfUpdate", "SpreadsheetTools",
                       "PresentationTools"):
            runtime._module(helper)
        with tempfile.TemporaryDirectory() as data:
            context = {"data_dir": data}
            workbook_spec = {"filename": "Quarterly.xlsx", "sheets": [{"name": "Summary", "rows": [["Quarter", "Revenue"], ["Q1", 120]]}]}
            workbook = runtime._save_artifact(context, json.dumps(workbook_spec), "create an Excel workbook")
            from openpyxl import load_workbook
            book = load_workbook(workbook, data_only=False)
            assert book["Summary"]["B2"].value == 120
            book.close()
            deck_spec = {"title": "Quarterly Review", "slides": [{"title": "Overview", "bullets": ["Revenue: 120"]}]}
            deck = runtime._save_artifact(context, json.dumps(deck_spec), "create a PowerPoint presentation")
            from pptx import Presentation
            presentation = Presentation(deck)
            assert len(presentation.slides) == 1
            assert any(shape.has_text_frame and "Overview" in shape.text for shape in presentation.slides[0].shapes)
            assert Path(deck).parent.resolve() == (Path(data) / "outputs").resolve()

    def test_refuses_output_inside_source_and_symlinked_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                package.package_source(ROOT, ROOT / "dist", "1.0")
            fake = Path(tmp) / "fake"
            fake.mkdir()
            for name in package._ROOT_FILES:
                (fake / name).write_text("x")
            (fake / "app").mkdir()
            for name in ("FileTools.py", "DeveloperAgent.py", "ProjectReadCache.py", "CADAgent.py",
                         "CADTools.py", "ProductSpecs.py", "MeshTools.py", "StockTools.py", "SelfUpdate.py",
                         "SpreadsheetTools.py", "PresentationTools.py"):
                (fake / "app" / name).write_text("x")
            (fake / "portable").mkdir()
            (fake / "portable/server.py").write_text("x")
            (fake / "portable/requirements.txt").write_text("x")
            target = Path(tmp) / "target.py"
            target.write_text("private")
            (fake / "portable/leak.py").symlink_to(target)
            with self.assertRaises(ValueError):
                package.collect_source(fake)


if __name__ == "__main__":
    unittest.main()
