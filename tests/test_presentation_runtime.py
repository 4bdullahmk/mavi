import importlib.util
import pathlib
import tempfile
import unittest

from portable import presentation_runtime as runtime


ROOT = pathlib.Path(__file__).resolve().parents[1]
HAS_PPTX = importlib.util.find_spec("pptx") is not None


class PresentationRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.output = pathlib.Path(self.temp.name) / "outputs"
        self.output.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def test_safe_basename_never_keeps_path_components(self):
        self.assertEqual(runtime.safe_output_basename("../Quarterly update!"), "Quarterly-update")
        self.assertEqual(runtime.safe_output_basename("東京"), "presentation")

    def test_rejects_invalid_specs_before_optional_dependency_load(self):
        invalid_specs = [
            {"title": "Deck", "slides": []},
            {"title": "Deck", "slides": [{"title": "Slide", "image_url": "https://example.test/a.png"}]},
            {"title": "Deck", "slides": [{"title": "Slide", "speaker_notes": "x" * (runtime.MAX_NOTES + 1)}]},
            {"title": "Deck", "slides": [{"title": "Slide"}] * (runtime.MAX_SLIDES + 1)},
        ]
        for spec in invalid_specs:
            with self.subTest(spec=spec):
                with self.assertRaises(ValueError):
                    runtime.create_presentation(spec, self.output)

    def test_rejects_output_folder_symlink(self):
        link = pathlib.Path(self.temp.name) / "linked-output"
        link.symlink_to(self.output, target_is_directory=True)
        with self.assertRaisesRegex(ValueError, "non-symlink"):
            runtime.create_presentation({"title": "Deck", "slides": [{"title": "One"}]}, link)

    @unittest.skipUnless(HAS_PPTX, "optional python-pptx dependency is not installed")
    def test_roundtrip_layout_content_notes_and_no_overwrite(self):
        spec = {
            "title": "Quarterly Review",
            "slides": [{
                "title": "Highlights",
                "body": "A concise overview.",
                "bullets": ["Revenue grew", "Costs stayed flat"],
                "speaker_notes": "Mention the regional comparison.",
            }],
        }
        first = runtime.create_presentation(spec, self.output)
        second = runtime.create_presentation(spec, self.output)
        self.assertEqual(first.name, "Quarterly-Review.pptx")
        self.assertEqual(second.name, "Quarterly-Review-2.pptx")
        self.assertTrue(first.is_file())
        self.assertTrue(second.is_file())

        from pptx import Presentation
        deck = Presentation(first)
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=3)
        self.assertEqual(deck.core_properties.title, "Quarterly Review")
        slide = deck.slides[0]
        text = [shape.text for shape in slide.shapes if shape.has_text_frame]
        combined_text = "\n".join(text)
        self.assertIn("Highlights", text)
        self.assertIn("A concise overview.", combined_text)
        self.assertIn("Revenue grew", combined_text)
        self.assertIn("Costs stayed flat", combined_text)
        self.assertEqual(slide.notes_slide.notes_text_frame.text, "Mention the regional comparison.")
        self.assertTrue(all(shape.shape_type != 13 for shape in slide.shapes))  # no picture assets

    @unittest.skipIf(HAS_PPTX, "only relevant when the optional dependency is absent")
    def test_missing_optional_dependency_has_actionable_error(self):
        with self.assertRaisesRegex(RuntimeError, "optional python-pptx"):
            runtime.create_presentation({"title": "Deck", "slides": [{"title": "One"}]}, self.output)

    def test_app_module_is_standalone_copy(self):
        path = ROOT / "app" / "PresentationTools.py"
        spec = importlib.util.spec_from_file_location("app_presentation_tools", path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        self.assertEqual(module.safe_output_basename("Same behavior"), runtime.safe_output_basename("Same behavior"))
        self.assertIsNot(module, runtime)


if __name__ == "__main__":
    unittest.main()
