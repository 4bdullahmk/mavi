import importlib.util
import builtins
import pathlib
import sys
import subprocess
import tempfile
import unittest
import json
from unittest.mock import patch

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "portable"))
import SpreadsheetTools as sheets
import FileTools
import tools_runtime


OPENPYXL = importlib.util.find_spec("openpyxl") is not None


def spec(rows=None, **sheet_fields):
    sheet = {"name": "Summary", "rows": rows if rows is not None else [["Month", "Sales"], ["Jan", 12]]}
    sheet.update(sheet_fields)
    return {"filename": "Quarterly Report.xlsx", "sheets": [sheet]}


class SpreadsheetRuntimeTests(unittest.TestCase):
    def test_validates_and_normalizes_explicit_cell_types(self):
        result = sheets.validate_spec(spec([["Name", "Total"], ["A", {"type": "formula", "value": "=SUM(B2:B2)"}]]))
        self.assertEqual(result["sheets"][0]["rows"][1][1], ("=SUM(B2:B2)", "formula"))
        self.assertIn("not evaluated", sheets.SPEC_GUIDE)

    def test_rejects_formula_network_external_macro_and_dde_vectors(self):
        dangerous = ["=HYPERLINK(\"https://evil.example\",\"open\")",
                     "=WEBSERVICE(\"https://evil.example\")",
                     "='[outside.xlsx]Sheet1'!A1", "=cmd|' /c calc'!A0",
                     "=CALL(\"calc.exe\")", "=RTD(\"server\",,\"x\")"]
        for formula in dangerous:
            with self.subTest(formula=formula), self.assertRaises(ValueError):
                sheets.validate_spec(spec([[{"type": "formula", "value": formula}]]))

    def test_formula_must_be_explicit_and_regular_strings_stay_strings(self):
        value = sheets.validate_spec(spec([["=SUM(A1:A3)"]]))
        self.assertEqual(value["sheets"][0]["rows"][0][0], ("=SUM(A1:A3)", "string"))
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([[{"type": "formula", "value": "SUM(A1:A3)"}]]))

    def test_enforces_sheet_row_column_spec_size_and_cell_limits(self):
        with self.assertRaises(ValueError):
            sheets.validate_spec({"filename": "x.xlsx", "sheets": []})
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([[] for _ in range(sheets.MAX_ROWS + 1)]))
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([[None] * (sheets.MAX_COLUMNS + 1)]))
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([["x" * 32768]]))
        huge = {"filename": "x.xlsx", "sheets": [{"name": "S", "rows": [["x" * 1_000_001]]}]}
        with self.assertRaises(ValueError):
            sheets.validate_spec(huge)

    def test_rejects_bad_sheet_table_range_chart_and_dropdown_references(self):
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([["A"], [1]], name="Bad/Name"))
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([["A", "A"], [1, 2]], tables=[{"name": "T", "range": "A1:B2"}]))
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([["A", "B"], [1, 2]], charts=[{"type": "bar", "title": "x", "data": "B1:B2", "categories": "A1:A2", "anchor": "D2"}]))
        with self.assertRaises(ValueError):
            sheets.validate_spec(spec([["Choice"]], validations=[{"range": "A1:A3", "type": "list", "formula1": "=Other!A1:A2"}]))

    def test_table_names_are_unique_across_workbook(self):
        value = {"filename": "x.xlsx", "sheets": [
            {"name": "One", "rows": [["A"], [1]], "tables": [{"name": "Shared", "range": "A1:A2"}]},
            {"name": "Two", "rows": [["A"], [1]], "tables": [{"name": "shared", "range": "A1:A2"}]}]}
        with self.assertRaises(ValueError):
            sheets.validate_spec(value)

    def test_create_only_accepts_outputs_directory(self):
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(ValueError):
                sheets.create_workbook(spec(), pathlib.Path(td) / "elsewhere")

    def test_dropdown_target_can_cover_bounded_empty_rows(self):
        result = sheets.validate_spec(spec([["Choice"], ["Open"]], validations=[
            {"range": "A2:A1000", "type": "list", "formula1": '"Open,Closed"'}]))
        self.assertEqual(result["sheets"][0]["validations"][0]["range"], "A2:A1000")

    def test_missing_openpyxl_has_actionable_error(self):
        original_import = builtins.__import__
        def without_openpyxl(name, *args, **kwargs):
            if name == "openpyxl" or name.startswith("openpyxl."):
                raise ImportError("simulated optional dependency missing")
            return original_import(name, *args, **kwargs)
        with patch("builtins.__import__", side_effect=without_openpyxl):
            with self.assertRaisesRegex(RuntimeError, "optional openpyxl"):
                sheets.build_workbook(spec())

    @unittest.skipUnless(OPENPYXL, "openpyxl is an optional runtime dependency not installed in this environment")
    def test_creates_formatted_workbook_with_charts_and_dropdowns(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            output = pathlib.Path(td) / "outputs"
            workbook_spec = spec(
                [["Month", "Sales"], ["Jan", 12], ["Feb", 15], ["Total", {"type": "formula", "value": "=SUM(B2:B3)"}]],
                tables=[{"name": "MonthlySales", "range": "A1:B3"}],
                charts=[{"type": "bar", "title": "Sales by month", "data": "B1:B3", "categories": "A2:A3", "anchor": "D2"}],
                number_formats=[{"range": "B2:B4", "format": "$#,##0.00"}],
                column_widths=[{"column": "A", "width": 18}],
                validations=[{"range": "A2:A4", "type": "list", "formula1": "\"Jan,Feb\""}])
            path = sheets.create_workbook(workbook_spec, output)
            book = load_workbook(path, data_only=False)
            ws = book["Summary"]
            self.assertEqual(ws["B4"].value, "=SUM(B2:B3)")
            self.assertEqual(ws["B2"].number_format, "$#,##0.00")
            self.assertEqual(len(ws.tables), 1)
            self.assertEqual(len(ws._charts), 1)
            self.assertEqual(len(ws.data_validations.dataValidation), 1)
            self.assertEqual(ws.freeze_panes, "A2")
            book.close()

    @unittest.skipUnless(OPENPYXL, "openpyxl is an optional runtime dependency not installed in this environment")
    def test_filetools_json_xlsx_roundtrip_uses_shared_builder(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            root = pathlib.Path(td)
            content = json.dumps(spec([["Item", "Count"], ["Widgets", 7],
                                       ["Total", {"type": "formula", "value": "=SUM(B2:B2)"}]]))
            result = FileTools.write({"root": str(root), "name": "Inventory.xlsx", "content": content})
            book = load_workbook(result["path"], data_only=False)
            self.assertEqual(book["Summary"]["B3"].value, "=SUM(B2:B2)")
            book.close()

    def test_filetools_schema_validation_is_side_effect_free(self):
        good_xlsx = spec([["Budget Item", "Amount"], ["Rent", 1200], ["Food", 300],
                          ["Total", {"type": "formula", "value": "=SUM(B2:B3)"}]],
                         charts=[{"type": "bar", "title": "Budget", "data": "B1:B3",
                                  "categories": "A2:A3", "anchor": "D2"}], freeze_panes="A2")
        bad_xlsx = spec([["Budget Item", "Amount"], ["Rent", 1200], ["Food", 300]],
                        charts=[{"type": "bar", "title": "Budget", "data": "B2:B3",
                                 "categories": "A2:A3", "anchor": "D2"}])
        with tempfile.TemporaryDirectory() as td:
            before = list(pathlib.Path(td).iterdir())
            self.assertEqual(FileTools.validate({"format": "xlsx", "content": json.dumps(good_xlsx)}), {"valid": True})
            bad_result = FileTools.validate({"format": "xlsx", "content": json.dumps(bad_xlsx)})
            self.assertFalse(bad_result["valid"])
            self.assertIn("Chart categories", bad_result["validation_error"])
            self.assertEqual(FileTools.validate({"format": "pptx", "content": json.dumps({"title": "Brief", "slides": [{"title": "Summary"}]})}), {"valid": True})
            self.assertFalse(FileTools.validate({"format": "pptx", "content": "not json"})["valid"])
            self.assertEqual(list(pathlib.Path(td).iterdir()), before)
        cli = subprocess.run([sys.executable, str(ROOT / "app" / "FileTools.py")],
                             input=json.dumps({"action":"validate","format":"xlsx","content":json.dumps(bad_xlsx)}),
                             text=True,capture_output=True,check=True)
        response=json.loads(cli.stdout)
        self.assertFalse(response["valid"])
        self.assertIn("Chart categories", response["validation_error"])
        self.assertNotIn("error", response)

    @unittest.skipUnless(OPENPYXL, "openpyxl is an optional runtime dependency not installed in this environment")
    def test_files_repairs_one_invalid_chart_range_and_keeps_requested_name(self):
        from openpyxl import load_workbook
        invalid = {"filename": "mavi-budget-check.xlsx", "sheets": [{"name": "Budget", "rows": [
            ["Budget Item", "Amount"], ["Rent", 1200], ["Food", 300],
            ["Total", {"type": "formula", "value": "=SUM(B2:B3)"}]],
            "charts": [{"type": "bar", "title": "Budget", "data": "B2:B3", "categories": "A2:A3", "anchor": "D2"}],
            "freeze_panes": "A2", "number_formats": [{"range": "B2:B4", "format": "$#,##0.00"}]}]}
        corrected = {"filename": "mavi-budget-check.xlsx", "sheets": [{"name": "Budget", "rows": [
            ["Budget Item", "Amount"], ["Rent", 1200], ["Food", 300],
            ["Total", {"type": "formula", "value": "=SUM(B2:B3)"}]],
            "charts": [{"type": "bar", "title": "Budget", "data": "B1:B3", "categories": "A2:A3", "anchor": "D2"}],
            "freeze_panes": "A2", "number_formats": [{"range": "B2:B4", "format": "$#,##0.00"}]}]}
        answers = iter([json.dumps(invalid), json.dumps(corrected)])
        calls = []
        context = {"model": "qwen3:8b", "progress": lambda _value: None,
                   "call_model": lambda messages, model=None: (calls.append(messages) or next(answers))}
        with tempfile.TemporaryDirectory() as td:
            context["data_dir"] = pathlib.Path(td)
            result = tools_runtime.run("files", "Create workbook mavi-budget-check.xlsx with a budget chart", [], context)
            output = pathlib.Path(td) / "outputs" / "mavi-budget-check.xlsx"
            self.assertIn("Saved locally", result)
            self.assertTrue(output.is_file())
            self.assertEqual(len(calls), 2)
            book = load_workbook(output, data_only=False)
            ws = book["Budget"]
            self.assertEqual(ws["B4"].value, "=SUM(B2:B3)")
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws["B2"].number_format, "$#,##0.00")
            self.assertEqual(len(ws._charts), 1)
            book.close()


if __name__ == "__main__":
    unittest.main()
