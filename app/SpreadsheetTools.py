"""Create bounded, formula-safe XLSX workbooks from declarative JSON specs.

Formulas are stored, never evaluated here. Excel or another spreadsheet app
recalculates them when the user opens the workbook.
"""
from __future__ import annotations

import json
import math
import os
import re
import tempfile
from pathlib import Path
from typing import Any

MAX_SPEC_BYTES = 1_000_000
MAX_SHEETS = 10
MAX_ROWS = 5_000
MAX_COLUMNS = 100
MAX_CELL_TEXT = 32_767
_SHEET_KEYS = {"name", "rows", "tables", "charts", "number_formats", "column_widths",
               "freeze_panes", "auto_filter", "validations"}
_FORMULA_BLOCK = re.compile(
    r"\b(?:HYPERLINK|WEBSERVICE|RTD|DDE|CMD|CALL|EXEC|REGISTER\.ID|EVALUATE)\b|"
    r"https?\s*:|ftp\s*:|www\.|[\[\]\|]", re.I)
_CELL_REF = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*)$", re.I)
_RANGE_REF = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*)(?::\$?([A-Z]{1,3})\$?([1-9][0-9]*))?$", re.I)
_TABLE_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]{0,254}$")
_FILE_STEM = re.compile(r"^[A-Za-z0-9][A-Za-z0-9 _.-]{0,63}$")


def schema_help() -> str:
    """Return concise authoring guidance suitable for a local model prompt."""
    example = {"filename": "Report.xlsx", "sheets": [{"name": "Summary",
        "rows": [["Month", "Sales"], ["Jan", 120]],
        "tables": [{"name": "SalesTable", "range": "A1:B2"}],
        "charts": [{"type": "bar", "title": "Sales", "data": "B1:B2",
                    "categories": "A2:A2", "anchor": "D2"}],
        "number_formats": [{"range": "B2:B2", "format": "$#,##0.00"}],
        "column_widths": [{"column": "A", "width": 18}], "freeze_panes": "A2",
        "auto_filter": True, "validations": [{"range": "A2:A2", "type": "list",
                                                  "formula1": '"Open,Closed"'}]}]}
    return (
        "Return one JSON object matching this example:\n" +
        json.dumps(example, ensure_ascii=False, indent=2) +
        "\nFormula cells must be objects {\"type\":\"formula\",\"value\":\"=SUM(B2:B10)\"}; "
        "formulas are saved but not evaluated until the workbook is opened in Excel. Use only "
        "internal references and ordinary calculation formulas. Chart data ranges include the "
        "header row (for example data B1:B3 with categories A2:A3); categories must cover exactly "
        "the data rows beneath that header. No macros, external links, URLs, "
        "DDE, command functions, HYPERLINK, or WEBSERVICE. Limits: 10 sheets, 5,000 rows and "
        "100 columns per sheet, and 1 MB JSON."
    )


def _column_number(label: str) -> int:
    number = 0
    for char in label.upper():
        if not "A" <= char <= "Z":
            raise ValueError("Invalid spreadsheet column reference.")
        number = number * 26 + ord(char) - 64
    return number


def _coordinate(value: Any, *, max_rows: int, max_columns: int, allow_edge: bool = False) -> tuple[int, int]:
    if not isinstance(value, str):
        raise ValueError("Expected an Excel cell reference such as A1.")
    match = _CELL_REF.fullmatch(value)
    if not match:
        raise ValueError("Invalid Excel cell reference.")
    col, row = _column_number(match.group(1)), int(match.group(2))
    row_limit = max_rows + 1 if allow_edge else max_rows
    col_limit = max_columns + 1 if allow_edge else max_columns
    if row > row_limit or col > col_limit:
        raise ValueError("Cell reference exceeds the populated sheet bounds.")
    return row, col


def _range(value: Any, *, max_rows: int, max_columns: int,
           min_rows: int = 1, min_columns: int = 1) -> tuple[int, int, int, int]:
    if not isinstance(value, str):
        raise ValueError("Expected a bounded Excel range such as A1:C12.")
    match = _RANGE_REF.fullmatch(value)
    if not match:
        raise ValueError("Invalid Excel range reference.")
    min_col, min_row = _column_number(match.group(1)), int(match.group(2))
    max_col = _column_number(match.group(3)) if match.group(3) else min_col
    max_row = int(match.group(4)) if match.group(4) else min_row
    if min_col > max_col or min_row > max_row:
        raise ValueError("Excel ranges must run from the top left to the bottom right.")
    if max_row > max_rows or max_col > max_columns:
        raise ValueError("Range exceeds the populated sheet bounds.")
    if max_row - min_row + 1 < min_rows or max_col - min_col + 1 < min_columns:
        raise ValueError("Range is too small for this workbook feature.")
    return min_col, min_row, max_col, max_row


def _check_formula(formula: str) -> str:
    if not isinstance(formula, str) or not formula.startswith("=") or len(formula) > 1_000:
        raise ValueError("Formula cells must contain an Excel formula of at most 1,000 characters.")
    if _FORMULA_BLOCK.search(formula):
        raise ValueError("External links, URLs, DDE, macros, and network or command formulas are not allowed.")
    return formula


def _cell_value(value: Any) -> tuple[Any, str]:
    if value is None:
        return None, "blank"
    if isinstance(value, bool):
        return value, "boolean"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("Cell numbers must be finite.")
        return value, "number"
    if isinstance(value, str):
        if len(value) > MAX_CELL_TEXT:
            raise ValueError("Cell text exceeds Excel's 32,767 character limit.")
        return value, "string"
    if not isinstance(value, dict) or set(value) != {"type", "value"}:
        raise ValueError("Typed cells must have exactly 'type' and 'value' fields.")
    kind, item = value["type"], value["value"]
    if kind == "formula":
        return _check_formula(item), "formula"
    if kind == "string" and isinstance(item, str) and len(item) <= MAX_CELL_TEXT:
        return item, "string"
    if kind == "number" and isinstance(item, (int, float)) and not isinstance(item, bool):
        if isinstance(item, float) and not math.isfinite(item):
            raise ValueError("Cell numbers must be finite.")
        return item, "number"
    if kind == "boolean" and isinstance(item, bool):
        return item, "boolean"
    if kind == "date" and isinstance(item, str):
        from datetime import date
        try:
            return date.fromisoformat(item), "date"
        except ValueError:
            pass
    raise ValueError("Typed cells support string, number, boolean, ISO date, and formula values.")


def validate_spec(spec: Any) -> dict:
    """Validate and normalize a workbook spec without importing openpyxl."""
    try:
        encoded = json.dumps(spec, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode("utf-8")
    except (TypeError, ValueError, UnicodeError, RecursionError):
        raise ValueError("Workbook spec must be valid JSON data.") from None
    if len(encoded) > MAX_SPEC_BYTES:
        raise ValueError("Workbook specification exceeds 1 MB.")
    if not isinstance(spec, dict) or set(spec) != {"filename", "sheets"}:
        raise ValueError("Workbook spec must contain only filename and sheets.")
    filename = spec["filename"]
    if (not isinstance(filename, str) or len(filename) > 80 or Path(filename).name != filename
            or "/" in filename or "\\" in filename or any(ord(c) < 32 for c in filename)
            or not filename.lower().endswith(".xlsx")):
        raise ValueError("Choose a plain .xlsx file name without folders.")
    stem = filename[:-5]
    if not _FILE_STEM.fullmatch(stem):
        raise ValueError("Workbook file name may use letters, numbers, spaces, dots, dashes, and underscores.")
    sheets = spec["sheets"]
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= MAX_SHEETS:
        raise ValueError("A workbook needs between 1 and 10 sheets.")
    clean_sheets, names, global_table_names = [], set(), set()
    for sheet in sheets:
        if not isinstance(sheet, dict) or set(sheet) - _SHEET_KEYS or not {"name", "rows"} <= set(sheet):
            raise ValueError("Each sheet needs a name and rows, with only supported formatting fields.")
        name = sheet["name"]
        if (not isinstance(name, str) or not name or len(name) > 31 or name[0] == "'" or name[-1] == "'"
                or any(c in name for c in "[]:*?/\\") or any(ord(c) < 32 for c in name)):
            raise ValueError("Sheet names must be 1–31 characters and cannot contain Excel-reserved characters.")
        if name.casefold() in names:
            raise ValueError("Sheet names must be unique, ignoring case.")
        names.add(name.casefold())
        rows = sheet["rows"]
        if not isinstance(rows, list) or len(rows) > MAX_ROWS:
            raise ValueError("Each sheet is limited to 5,000 rows.")
        if any(not isinstance(row, list) or len(row) > MAX_COLUMNS for row in rows):
            raise ValueError("Each sheet is limited to 100 columns.")
        clean_rows = []
        max_columns = max((len(row) for row in rows), default=0)
        for row in rows:
            clean_rows.append([_cell_value(value) for value in row])
        max_rows = len(rows)
        clean = {"name": name, "rows": clean_rows, "max_rows": max_rows,
                 "max_columns": max_columns}

        tables = sheet.get("tables", [])
        if not isinstance(tables, list) or len(tables) > 20:
            raise ValueError("Each sheet can have up to 20 formatted tables.")
        clean["tables"] = []
        occupied = []
        for table in tables:
            if not isinstance(table, dict) or set(table) - {"name", "range", "style"} or not {"name", "range"} <= set(table):
                raise ValueError("A table needs a name and range.")
            tname, tref = table["name"], table["range"]
            if not isinstance(tname, str) or not _TABLE_NAME.fullmatch(tname) or re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]*", tname):
                raise ValueError("Invalid Excel table name.")
            if tname.casefold() in global_table_names:
                raise ValueError("Table names must be unique in a workbook.")
            global_table_names.add(tname.casefold())
            bounds = _range(tref, max_rows=max_rows, max_columns=max_columns, min_rows=2, min_columns=1)
            if any(not (bounds[2] < old[0] or bounds[0] > old[2] or bounds[3] < old[1] or bounds[1] > old[3]) for old in occupied):
                raise ValueError("Formatted table ranges cannot overlap.")
            occupied.append(bounds)
            headers = [clean_rows[bounds[1] - 1][col - 1] for col in range(bounds[0], bounds[2] + 1)]
            if any(kind != "string" or not isinstance(value, str) or not value.strip() for value, kind in headers):
                raise ValueError("Formatted table headers must be nonempty text.")
            if len({value.casefold() for value, _ in headers}) != len(headers):
                raise ValueError("Formatted table headers must be unique.")
            style = table.get("style", "TableStyleMedium2")
            if style not in ("TableStyleMedium2", "TableStyleMedium4", "TableStyleMedium9", "TableStyleLight1", "TableStyleLight9"):
                raise ValueError("Choose a supported built-in table style.")
            clean["tables"].append({"name": tname, "range": tref, "style": style})

        formats = sheet.get("number_formats", [])
        if not isinstance(formats, list) or len(formats) > 100:
            raise ValueError("Each sheet can have up to 100 number-format ranges.")
        clean["number_formats"] = []
        for fmt in formats:
            if not isinstance(fmt, dict) or set(fmt) != {"range", "format"} or not isinstance(fmt["format"], str) or not 1 <= len(fmt["format"]) <= 100:
                raise ValueError("Number formats need a range and a short format string.")
            _range(fmt["range"], max_rows=max_rows, max_columns=max_columns)
            clean["number_formats"].append(dict(fmt))

        widths = sheet.get("column_widths", [])
        if not isinstance(widths, list) or len(widths) > MAX_COLUMNS:
            raise ValueError("Column widths must be a list of at most 100 columns.")
        clean["column_widths"] = []
        seen_widths = set()
        for width in widths:
            if not isinstance(width, dict) or set(width) != {"column", "width"}:
                raise ValueError("Column widths need column and width fields.")
            col = width["column"]
            if not isinstance(col, str) or not re.fullmatch(r"[A-Z]{1,3}", col.upper()):
                raise ValueError("Column width needs a column label such as A or BC.")
            col_num = _column_number(col)
            value = width["width"]
            if col_num > max(MAX_COLUMNS, max_columns) or isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 4 <= value <= 80:
                raise ValueError("Column widths must be between 4 and 80 characters and within the 100-column limit.")
            if col_num in seen_widths:
                raise ValueError("Column widths must not repeat a column.")
            seen_widths.add(col_num)
            clean["column_widths"].append((col_num, float(value)))

        freeze = sheet.get("freeze_panes")
        if freeze is not None:
            row, col = _coordinate(freeze, max_rows=max_rows, max_columns=max_columns, allow_edge=True)
            if row < 2 and col < 2:
                raise ValueError("Freeze panes must leave at least one row or column frozen.")
        clean["freeze_panes"] = freeze
        auto_filter = sheet.get("auto_filter", True)
        if not isinstance(auto_filter, bool):
            raise ValueError("auto_filter must be true or false.")
        clean["auto_filter"] = auto_filter

        validations = sheet.get("validations", [])
        if not isinstance(validations, list) or len(validations) > 100:
            raise ValueError("Each sheet can have up to 100 dropdown ranges.")
        clean["validations"] = []
        for validation in validations:
            if (not isinstance(validation, dict) or set(validation) - {"range", "type", "formula1", "allow_blank"}
                    or not {"range", "type", "formula1"} <= set(validation)):
                raise ValueError("Dropdown validation needs range, type, and formula1 fields.")
            if validation["type"] != "list":
                raise ValueError("Only list dropdown validation is supported.")
            # Dropdown targets may intentionally cover blank future-entry cells.
            _range(validation["range"], max_rows=MAX_ROWS, max_columns=MAX_COLUMNS)
            formula1 = validation["formula1"]
            if not isinstance(formula1, str) or len(formula1) > 255 or any(ord(c) < 32 for c in formula1):
                raise ValueError("Dropdown source must be a short list or a bounded same-sheet range.")
            list_match = re.fullmatch(r'"[^"\r\n]{1,253}"', formula1)
            range_value = formula1[1:] if formula1.startswith("=") else formula1
            range_match = _RANGE_REF.fullmatch(range_value)
            if list_match:
                pass
            elif range_match:
                _range(range_value, max_rows=max_rows, max_columns=max_columns)
                formula1 = "=" + range_value
            else:
                raise ValueError("Dropdown source must be a comma-delimited quoted list or a same-sheet range.")
            allow_blank = validation.get("allow_blank", True)
            if not isinstance(allow_blank, bool):
                raise ValueError("allow_blank must be true or false.")
            clean["validations"].append({"range": validation["range"], "formula1": formula1,
                                         "allow_blank": allow_blank})

        charts = sheet.get("charts", [])
        if not isinstance(charts, list) or len(charts) > 10:
            raise ValueError("Each sheet can have up to 10 charts.")
        clean["charts"] = []
        for chart in charts:
            if not isinstance(chart, dict) or set(chart) != {"type", "title", "data", "categories", "anchor"}:
                raise ValueError("Charts need type, title, data, categories, and anchor fields.")
            if chart["type"] not in ("bar", "line") or not isinstance(chart["title"], str) or not 1 <= len(chart["title"]) <= 100:
                raise ValueError("Charts support titled bar and line charts.")
            data_bounds = _range(chart["data"], max_rows=max_rows, max_columns=max_columns, min_rows=2)
            category_bounds = _range(chart["categories"], max_rows=max_rows, max_columns=max_columns)
            if category_bounds[0] != category_bounds[2] or category_bounds[1] != data_bounds[1] + 1 or category_bounds[3] != data_bounds[3]:
                raise ValueError("Chart categories must be one column matching the data rows below its header.")
            _coordinate(chart["anchor"], max_rows=MAX_ROWS, max_columns=MAX_COLUMNS)
            clean["charts"].append(dict(chart))
        clean_sheets.append(clean)
    return {"filename": filename, "sheets": clean_sheets}


def _build_validated(clean: dict):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils.cell import get_column_letter
    except ImportError:
        raise RuntimeError("Excel workbook creation needs the optional openpyxl package.") from None

    book = Workbook()
    book.remove(book.active)
    for item in clean["sheets"]:
        sheet = book.create_sheet(item["name"])
        for row_number, row in enumerate(item["rows"], start=1):
            for col_number, (value, kind) in enumerate(row, start=1):
                cell = sheet.cell(row_number, col_number, value)
                if kind == "string":
                    cell.data_type = "s"  # A leading '=' remains literal text unless typed as formula.
        if item["max_rows"] and item["max_columns"]:
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 24
            sheet.freeze_panes = item["freeze_panes"] or "A2"
            if item["auto_filter"]:
                sheet.auto_filter.ref = f"A1:{get_column_letter(item['max_columns'])}{item['max_rows']}"
        for table in item["tables"]:
            obj = Table(displayName=table["name"], ref=table["range"])
            obj.tableStyleInfo = TableStyleInfo(name=table["style"], showFirstColumn=False,
                                                showLastColumn=False, showRowStripes=True,
                                                showColumnStripes=False)
            sheet.add_table(obj)
        for fmt in item["number_formats"]:
            c1, r1, c2, r2 = _range(fmt["range"], max_rows=item["max_rows"], max_columns=item["max_columns"])
            for row in sheet.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                for cell in row:
                    cell.number_format = fmt["format"]
        for col, width in item["column_widths"]:
            sheet.column_dimensions[get_column_letter(col)].width = width
        for validation in item["validations"]:
            rule = DataValidation(type="list", formula1=validation["formula1"],
                                  allow_blank=validation["allow_blank"])
            rule.error = "Choose a value from the list."
            rule.errorTitle = "Invalid choice"
            rule.showErrorMessage = True
            sheet.add_data_validation(rule)
            rule.add(validation["range"])
        for chart_spec in item["charts"]:
            chart = BarChart() if chart_spec["type"] == "bar" else LineChart()
            if chart_spec["type"] == "bar":
                chart.type = "col"
                chart.grouping = "clustered"
            chart.title = chart_spec["title"]
            chart.style = 10
            c1, r1, c2, r2 = _range(chart_spec["data"], max_rows=item["max_rows"], max_columns=item["max_columns"], min_rows=2)
            cat_col, cat_start, _, cat_end = _range(chart_spec["categories"], max_rows=item["max_rows"], max_columns=item["max_columns"])
            chart.add_data(Reference(sheet, min_col=c1, max_col=c2, min_row=r1, max_row=r2), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=cat_col, min_row=cat_start, max_row=cat_end))
            chart.height, chart.width = 8, 14
            sheet.add_chart(chart, chart_spec["anchor"])

    # Excel recalculates stored formulas on open. Mavi never evaluates them.
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.calculation.calcMode = "auto"
    return book


SPEC_GUIDE = schema_help()


def build_workbook(spec: Any):
    """Build a validated workbook in memory for native or portable file writers."""
    return _build_validated(validate_spec(spec))


def create_workbook(spec: Any, output_dir: str | Path) -> Path:
    """Create a new unique .xlsx file inside the caller's DATA/outputs folder."""
    clean = validate_spec(spec)
    raw_dir = Path(output_dir).absolute()
    if raw_dir.name.casefold() != "outputs" or raw_dir.is_symlink():
        raise ValueError("Workbook files can only be saved in the local outputs folder.")
    raw_dir.mkdir(parents=True, exist_ok=True)
    if raw_dir.is_symlink():
        raise ValueError("The outputs folder cannot be a symbolic link.")
    directory = raw_dir.resolve(strict=True)
    book = _build_validated(clean)
    stem = clean["filename"][:-5].strip(" .") or "Workbook"
    if not _FILE_STEM.fullmatch(stem):
        stem = "Workbook"
    fd, temp_name = tempfile.mkstemp(prefix=stem + "-", suffix=".xlsx", dir=directory)
    os.close(fd)
    target = Path(temp_name)
    try:
        book.save(target)
        return target
    except Exception:
        target.unlink(missing_ok=True)
        raise
